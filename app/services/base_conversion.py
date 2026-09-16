"""El mismo motor del script, con registro transaccional de sus resultados."""
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from app.database import Database
from scripts.transformar_base import cell_text, convert_files, normalize


def normalize_phone_number(value) -> str:
    digits = re.sub(r"\D", "", cell_text(value))
    if digits.startswith("593") and len(digits) == 12:
        digits = "0" + digits[3:]
    elif len(digits) == 9 and digits.startswith("9"):
        digits = "0" + digits
    return digits


def normalize_call_date(value) -> str:
    """Convierte la fecha de Hoja1 a AAAAMMDD, como aparece en el audio."""
    if isinstance(value, datetime):
        return value.strftime("%Y%m%d")
    if isinstance(value, date):
        return value.strftime("%Y%m%d")
    text = cell_text(value).strip()
    for pattern in (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y%m%d%H%M%S", "%Y%m%d",
    ):
        try:
            return datetime.strptime(text, pattern).strftime("%Y%m%d")
        except ValueError:
            continue
    match = re.search(r"\b(20\d{2})[-/]?(0[1-9]|1[0-2])[-/]?([0-2]\d|3[01])\b", text)
    return "".join(match.groups()) if match else ""


@dataclass(frozen=True)
class BaseAudioIndex:
    phone_dates: dict[str, frozenset[str]]
    names: dict[str, str] = field(default_factory=dict)

    @property
    def phones(self) -> set[str]:
        return set(self.phone_dates)

    @property
    def dates(self) -> set[str]:
        return {day for days in self.phone_dates.values() for day in days}


def load_hoja1_audio_index(path: Path) -> BaseAudioIndex:
    """Lee teléfonos, fechas y nombres de Hoja1 para relacionarlos con los audios."""
    path = Path(path).resolve()
    if not path.is_file() or path.suffix.casefold() != ".xlsx":
        raise ValueError("Selecciona un Excel transformado que todavía exista.")
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException, ParseError) as exc:
        raise ValueError(f"No se pudo leer el Excel transformado: {exc}") from exc
    try:
        if "Hoja1" not in book.sheetnames:
            raise ValueError("El Excel seleccionado no contiene la hoja Hoja1.")
        sheet = book["Hoja1"]
        if normalize(sheet.cell(1, 1).value) != "telefono":
            raise ValueError("Hoja1 no tiene la columna Teléfono en la primera posición.")
        phone_dates: dict[str, set[str]] = {}
        names: dict[str, str] = {}
        for phone_value, _state, _agent, date_value, _duration, name_value in sheet.iter_rows(
            min_row=2, min_col=1, max_col=6, values_only=True
        ):
            phone = normalize_phone_number(phone_value)
            if not phone:
                continue
            phone_dates.setdefault(phone, set())
            name = cell_text(name_value).strip()
            if name and phone not in names:
                names[phone] = name
            day = normalize_call_date(date_value)
            if day:
                phone_dates[phone].add(day)
        return BaseAudioIndex(
            {phone: frozenset(days) for phone, days in phone_dates.items()},
            names,
        )
    finally:
        book.close()


def load_hoja1_phones(path: Path) -> set[str]:
    """Lee como identificadores los teléfonos únicos de Hoja1."""
    return load_hoja1_audio_index(path).phones


def process_bases(database: Database, paths: list[Path], output_folder: Path, **options):
    job_id = database.start_job(paths, {**options, "output_folder": str(output_folder)})
    result = None
    try:
        result = convert_files(paths, output_folder=output_folder, **options)
        book = openpyxl.load_workbook(result.output, read_only=True, data_only=True)
        try:
            database.finish_job(job_id, result, book["Hoja1"].iter_rows(min_row=2, max_col=11, values_only=True))
        finally:
            book.close()
        return result
    except Exception as exc:
        # Un Excel ya publicado permanece disponible aunque falle la persistencia.
        try:
            database.fail_job(job_id, exc, result.output if result else None)
        except Exception:
            pass  # Conserva el error original, no un segundo error de SQLite.
        if result is not None:
            raise RuntimeError(f"El Excel está guardado en {result.output}, pero falló el historial: {exc}") from exc
        raise
