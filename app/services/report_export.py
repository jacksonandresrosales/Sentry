"""Exportación de denuncias relacionadas con la base activa de Sentry."""
from __future__ import annotations

from dataclasses import dataclass
import os
from pathlib import Path
import tempfile
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
import xlsxwriter

from app.services.base_conversion import hoja1_column_layout, normalize_phone_number
from scripts.transformar_base import DEFAULT_STATE, LUCID_HEADERS, cell_text


EXPORT_MODE_LABELS = {
    "automatic": "denuncias automáticas sin verificar",
    "verified": "denuncias verificadas",
    "complaints": "todas las denuncias",
    "all": "toda la base",
}


@dataclass(frozen=True)
class BaseExportRecord:
    phone: str
    customer: str
    identifier: str
    state: str = DEFAULT_STATE
    source_system: str = "issabel"


@dataclass(frozen=True)
class ExportResult:
    output: Path
    row_count: int
    complaint_count: int


def load_base_export_records(path: Path) -> list[BaseExportRecord]:
    """Lee Hoja1 en el formato de origen y conserva la primera fila por teléfono."""
    path = Path(path).resolve()
    if not path.is_file() or path.suffix.casefold() != ".xlsx":
        raise ValueError("La base activa no está disponible. Selecciona nuevamente el Excel transformado.")
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException, ParseError) as exc:
        raise ValueError(f"No se pudo leer la base activa: {exc}") from exc
    try:
        if "Hoja1" not in book.sheetnames:
            raise ValueError("La base activa no contiene la hoja Hoja1.")
        sheet = book["Hoja1"]
        source_system, headers = hoja1_column_layout(sheet)
        required = (("telefono", "nombre", "estado") if source_system == "lucid"
                    else ("telefono", "nombre", "id"))
        if "id" in headers and "id" not in required:
            required += ("id",)

        records: list[BaseExportRecord] = []
        seen: set[str] = set()
        maximum = max(headers[key] for key in required) + 1
        for values in sheet.iter_rows(min_row=2, max_col=maximum, values_only=True):
            phone = normalize_phone_number(values[headers["telefono"]])
            if not phone or phone in seen:
                continue
            seen.add(phone)
            records.append(BaseExportRecord(
                phone=phone,
                customer=cell_text(values[headers["nombre"]]),
                identifier=cell_text(values[headers["id"]]) if "id" in headers else "",
                state=cell_text(values[headers["estado"]]) if source_system == "lucid" else DEFAULT_STATE,
                source_system=source_system,
            ))
        return records
    finally:
        book.close()


def select_export_records(
    records: list[BaseExportRecord],
    automatic_phones: set[str],
    verified_phones: set[str],
    mode: str,
) -> list[tuple[BaseExportRecord, bool]]:
    """Selecciona filas y devuelve si cada una debe resaltarse como denuncia."""
    if mode not in EXPORT_MODE_LABELS:
        raise ValueError("Opción de exportación desconocida.")
    automatic = {phone for value in automatic_phones if (phone := normalize_phone_number(value))}
    verified = {phone for value in verified_phones if (phone := normalize_phone_number(value))}
    complaint = automatic | verified
    automatic_unverified = automatic - verified
    if mode == "automatic":
        selected = automatic_unverified
    elif mode == "verified":
        selected = verified
    elif mode == "complaints":
        selected = complaint
    else:
        selected = None
    return [
        (record, record.phone in complaint)
        for record in records
        if selected is None or record.phone in selected
    ]


def export_audit_excel(
    output: Path,
    base_path: Path,
    automatic_phones: set[str],
    verified_phones: set[str],
    mode: str,
) -> ExportResult:
    """Exporta las columnas de auditoría correspondientes al sistema de la base."""
    output = Path(output).resolve()
    if output.suffix.casefold() != ".xlsx":
        output = output.with_suffix(".xlsx")
    base_path = Path(base_path).resolve()
    if output == base_path or (output.exists() and base_path.exists() and output.samefile(base_path)):
        raise ValueError("La exportación no puede reemplazar la base activa. Selecciona otro archivo de salida.")
    records = load_base_export_records(base_path)
    selected = select_export_records(records, automatic_phones, verified_phones, mode)
    if not selected:
        raise ValueError(f"No hay {EXPORT_MODE_LABELS[mode]} para exportar.")

    source_system = selected[0][0].source_system
    verified = {phone for value in verified_phones if (phone := normalize_phone_number(value))}
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=".sentry-export-", suffix=".xlsx", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with xlsxwriter.Workbook(str(temporary)) as book:
            book.set_properties({"title": "Exportación de auditoría Sentry"})
            sheet = book.add_worksheet("Resultados")
            header = book.add_format({
                "font_name": "Aptos", "font_size": 11, "bold": True,
                "font_color": "#FFFFFF", "bg_color": "#202220",
                "align": "center", "valign": "vcenter",
            })
            complaint_format = book.add_format({
                "font_name": "Aptos", "font_size": 11, "font_color": "#9C0006",
                "bg_color": "#FFC7CE", "num_format": "@", "valign": "vcenter",
            })
            normal_format = book.add_format({
                "font_name": "Aptos", "font_size": 11, "font_color": "#006100",
                "bg_color": "#C6EFCE", "num_format": "@", "valign": "vcenter",
            })
            headers = (LUCID_HEADERS if source_system == "lucid" else
                       ("Número de celular", "Nombre del cliente", "ID", "Estado"))
            for column, title in enumerate(headers):
                sheet.write_string(0, column, title, header)
            complaint_count = 0
            for row_index, (record, is_complaint) in enumerate(selected, 1):
                row_format = complaint_format if is_complaint else normal_format
                complaint_count += int(is_complaint)
                if source_system == "lucid":
                    state = ("Denuncia verificada" if record.phone in verified else
                             "Denuncia detectada" if is_complaint else record.state)
                    values = (record.phone, record.customer, record.identifier, state)
                else:
                    values = (record.phone, record.customer, record.identifier, DEFAULT_STATE)
                for column, value in enumerate(values):
                    sheet.write_string(row_index, column, value, row_format)
            sheet.set_column(0, 0, 20)
            sheet.set_column(1, 1, 42)
            if source_system == "lucid":
                sheet.set_column(2, 2, 22)
                sheet.set_column(3, 3, 36)
            else:
                sheet.set_column(2, 2, 22)
                sheet.set_column(3, 3, 32)
            sheet.set_row(0, 22)
            sheet.freeze_panes(1, 0)
            sheet.autofilter(0, 0, len(selected), len(headers) - 1)
        temporary.replace(output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return ExportResult(output, len(selected), complaint_count)
