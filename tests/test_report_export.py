from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from app.services.report_export import export_audit_excel
from scripts.transformar_base import DEFAULT_STATE, HEADERS


class ReportExportTests(unittest.TestCase):
    def make_base(self, path: Path) -> None:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Hoja1"
        sheet.append(HEADERS)
        sheet.append(["0990000001", "Success", "001", "2026-09-15", "30",
                      "CLIENTE AUTOMATICO", "0123456789", "CÉDULA", DEFAULT_STATE, "V", "B1"])
        sheet.append(["0980000002", "Success", "002", "2026-09-15", "40",
                      "CLIENTE VERIFICADO", "0987654321", "CÉDULA", DEFAULT_STATE, "V", "B1"])
        sheet.append(["0970000003", "Success", "003", "2026-09-15", "50",
                      "CLIENTE NORMAL", "0912345678", "CÉDULA", DEFAULT_STATE, "V", "B1"])
        book.save(path)
        book.close()

    def test_exports_four_columns_and_colors_complaints_red(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            base = folder / "base.xlsx"
            output = folder / "completa.xlsx"
            self.make_base(base)

            result = export_audit_excel(
                output,
                base,
                {"0990000001", "0980000002"},
                {"0980000002"},
                "all",
            )

            self.assertEqual((result.row_count, result.complaint_count), (3, 2))
            book = openpyxl.load_workbook(output, data_only=True)
            try:
                sheet = book["Resultados"]
                self.assertEqual(
                    [cell.value for cell in sheet[1]],
                    ["Número de celular", "Nombre del cliente", "ID", "Estado"],
                )
                self.assertEqual(sheet.max_column, 4)
                self.assertEqual(sheet["D2"].value, DEFAULT_STATE)
                self.assertTrue(sheet["A2"].fill.fgColor.rgb.endswith("FFC7CE"))
                self.assertTrue(sheet["A3"].fill.fgColor.rgb.endswith("FFC7CE"))
                self.assertTrue(sheet["A4"].fill.fgColor.rgb.endswith("C6EFCE"))
            finally:
                book.close()

    def test_automatic_and_verified_exports_are_separate(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            base = folder / "base.xlsx"
            self.make_base(base)
            automatic = export_audit_excel(
                folder / "automaticas.xlsx", base,
                {"0990000001", "0980000002"}, {"0980000002"}, "automatic",
            )
            verified = export_audit_excel(
                folder / "verificadas.xlsx", base,
                {"0990000001", "0980000002"}, {"0980000002"}, "verified",
            )
            self.assertEqual((automatic.row_count, verified.row_count), (1, 1))
            automatic_book = openpyxl.load_workbook(automatic.output, read_only=True, data_only=True)
            verified_book = openpyxl.load_workbook(verified.output, read_only=True, data_only=True)
            try:
                self.assertEqual(automatic_book["Resultados"]["A2"].value, "0990000001")
                self.assertEqual(verified_book["Resultados"]["A2"].value, "0980000002")
            finally:
                automatic_book.close()
                verified_book.close()


if __name__ == "__main__":
    unittest.main()
