"""Pruebas de conversión completas con datos ficticios, sin datos de clientes."""

import csv
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from scripts.transformar_base import (
    BaseError, OutputCollisionError, DEFAULT_STATE, HEADERS, convert_files, default_output, main,
    read_csv, read_xlsx, transform, write_workbook,
)


RAW_HEADERS = [
    "Teléfono", "Estado Llamada", "Agente", "Fecha y hora", "Duración(Seg)",
    "Uniqueid", "Failure Code", "Failure Cause", "Nombre", "ID", "ESTADO",
]


def record(phone="0990000001", identifier="0123456789", state=DEFAULT_STATE,
           call_state="Success", name="PERSONA DE PRUEBA"):
    return [phone, call_state, "00123", "2026-07-01 12:00:00", "37", "call-1", "", "",
            name, identifier, state]


class TransformTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def csv_file(self, rows, headers=RAW_HEADERS, name="campana B1.csv",
                 delimiter=",", encoding="utf-8-sig"):
        path = self.folder / name
        with path.open("w", encoding=encoding, newline="") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            writer.writerow([""] * (len(headers) - 1) + ["FORMULARIO"])
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def test_conversion_filters_deduplicates_and_caches_real_formulas(self):
        path = self.csv_file([
            record(), record(name="SEGUNDA LLAMADA"),
            record(phone="0990000002", identifier="0123456789001"),
            record(phone="0990000003", identifier="AB12345"),
            record(phone="0990000004", state="BUZON"),
            record(phone="0990000005", call_state="ShortCall"),
        ])
        source = read_csv(path, None, None)[0]
        selected = transform(source, base_type="V")
        output = self.folder / "salida.xlsx"
        self.assertEqual(write_workbook(output, [source], [selected]), (4, 3))
        values = openpyxl.load_workbook(output, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(output, read_only=True, data_only=False)
        try:
            self.assertEqual(values.sheetnames, ["campana B1", "BASE1", "Validar", "Hoja1"])
            sheet = values["Hoja1"]
            self.assertEqual(tuple(c.value for c in sheet[1]), HEADERS)
            self.assertEqual(sheet["A2"].value, "0990000001")
            self.assertEqual(sheet["C2"].value, "00123")
            self.assertEqual(sheet["G2"].value, "0123456789")
            self.assertEqual(sheet["H2"].value, "CÉDULA")
            self.assertEqual(sheet["H3"].value, "RUC")
            self.assertIn(sheet["H4"].value, ("", None))
            self.assertEqual(sheet["F2"].value, "PERSONA DE PRUEBA")
            self.assertEqual(sheet["J2"].value, "V")
            self.assertEqual(sheet["K2"].value, "B1")
            self.assertEqual(formulas["Validar"]["H2"].value,
                             '=IF(LEN(G2)=10,"CÉDULA",IF(LEN(G2)=13,"RUC",""))')
            self.assertIn("UNIQUE(Validar!A1:A5)", formulas["Hoja1"]["A1"].value.text)
            self.assertEqual(formulas["Hoja1"]["A1"].value.ref, "A1:A4")
            self.assertIn("VLOOKUP(_xlfn.ANCHORARRAY($A$1)", formulas["Hoja1"]["H1"].value.text)
            self.assertEqual(sheet["G2"].data_type, "s")
        finally:
            values.close()
            formulas.close()

    def test_damaged_headers_and_default_verification_mark(self):
        headers = [h.replace("é", "�").replace("ó", "�") for h in RAW_HEADERS]
        source = read_csv(self.csv_file([record()], headers), None, None)[0]
        self.assertEqual(transform(source)[0][7:11], ["CÉDULA", DEFAULT_STATE, "V", "B1"])

    def test_semicolon_cp1252_and_quoted_multiline_values(self):
        path = self.csv_file([record(name="PRUEBA; ÁREA\nDOS")],
                             delimiter=";", encoding="cp1252")
        source = read_csv(path, None, None)[0]
        self.assertEqual(source.rows[0][8], "PRUEBA; ÁREA\nDOS")

    def test_missing_and_malformed_columns_are_rejected(self):
        path = self.csv_file([record()[:-1]], RAW_HEADERS[:-1])
        with self.assertRaisesRegex(BaseError, "ESTADO"):
            read_csv(path, None, None)
        path = self.csv_file([record() + ["extra"]])
        with self.assertRaisesRegex(BaseError, "columnas"):
            read_csv(path, None, ",")

    def test_configurable_filters_and_metadata(self):
        headers = RAW_HEADERS + ["T BASE", "N° BASE"]
        path = self.csv_file([record(call_state="ShortCall") + ["R", "B9"]], headers)
        source = read_csv(path, None, None)[0]
        self.assertEqual(transform(source), [])
        self.assertEqual(transform(source, call_state="*")[0][9:11], ["R", "B9"])
        self.assertEqual(transform(source, call_state="*", base_type="V", base_number="B4")
                         [0][9:11], ["V", "B4"])

    def test_no_input_overwrite_and_no_accidental_output_overwrite(self):
        path = self.csv_file([record()])
        source = read_csv(path, None, None)[0]
        output = self.folder / "salida.xlsx"
        write_workbook(output, [source], [transform(source)])
        before = output.read_bytes()
        with self.assertRaises(BaseError):
            write_workbook(output, [source], [[]])
        self.assertEqual(output.read_bytes(), before)
        self.assertEqual(main([str(output), "-o", str(output), "--sobrescribir"]), 1)
        self.assertEqual(output.read_bytes(), before)

    def test_data_starting_with_equal_sign_remains_text(self):
        source = read_csv(self.csv_file([record(name='=HYPERLINK("https://example.com")')]),
                          None, None)[0]
        output = self.folder / "salida.xlsx"
        write_workbook(output, [source], [transform(source)])
        book = openpyxl.load_workbook(output, read_only=True, data_only=False)
        try:
            self.assertEqual(book["Hoja1"]["F2"].data_type, "s")
            self.assertEqual(book["Hoja1"]["F2"].value, '=HYPERLINK("https://example.com")')
        finally:
            book.close()

    def test_xlsx_recovers_formatted_zeros_and_ignores_derived_sheets(self):
        path = self.folder / "origen.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "campana B2"
        sheet.append([f"Column{i}" for i in range(1, len(RAW_HEADERS) + 1)])
        sheet.append(RAW_HEADERS)
        sheet.append(record(phone=990000001, identifier=123456789))
        sheet["A3"].number_format = "0000000000"
        sheet["J3"].number_format = "0000000000"
        derived = book.create_sheet("Validar")
        derived.append(HEADERS)
        derived.append(["0990000001", "Success", "00123", "2026-07-01 12:00:00", "37",
                        "PERSONA DE PRUEBA", "0123456789", "CÉDULA", DEFAULT_STATE, "R", "B2"])
        book.save(path)
        book.close()
        sources = read_xlsx(path, None)
        self.assertEqual(len(sources), 1)
        self.assertEqual(transform(sources[0])[0][0], "0990000001")
        self.assertEqual(transform(sources[0])[0][6], "0123456789")
        self.assertEqual(read_xlsx(path, "Validar")[0].name, "Validar")

    def test_multi_base_cli_consolidates_first_phone_occurrence(self):
        first = self.csv_file([record()], name="campana B1.csv")
        second = self.csv_file([record(name="OTRA BASE"), record(phone="0990000002")],
                               name="campana B2.csv")
        output = self.folder / "consolidado.xlsx"
        self.assertEqual(main([str(first), str(second), "-o", str(output)]), 0)
        book = openpyxl.load_workbook(output, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Validar"].max_row, 4)
            self.assertEqual(book["Hoja1"].max_row, 3)
            self.assertEqual(book["Hoja1"]["K2"].value, "B1")
            self.assertEqual(book["Hoja1"]["K3"].value, "B2")
        finally:
            book.close()

    def test_empty_filter_results_still_create_headers(self):
        source = read_csv(self.csv_file([record(state="BUZON")]), None, None)[0]
        output = self.folder / "vacio.xlsx"
        self.assertEqual(write_workbook(output, [source], [transform(source)]), (0, 0))
        book = openpyxl.load_workbook(output, read_only=True)
        try:
            self.assertEqual(book["Hoja1"].max_row, 1)
            self.assertEqual(book[source.name].max_row, 4)
        finally:
            book.close()

    def test_corrupt_xlsx_is_reported_without_creating_output(self):
        path = self.folder / "corrupto.xlsx"
        path.write_bytes(b"not an Excel file")
        output = self.folder / "salida.xlsx"
        self.assertEqual(main([str(path), "-o", str(output)]), 1)
        self.assertFalse(output.exists())

    def test_db_delete_filename_and_shared_conversion_api(self):
        path = self.csv_file([record()], name="BASE original B1.csv")
        output = default_output(path, self.folder)
        self.assertEqual(output.name, "BASE original B1_DB_delete_B1.xlsx")
        result = convert_files([path], output)
        self.assertEqual(result.output, output.resolve())
        self.assertEqual((result.filtered, result.unique), (1, 1))
        with self.assertRaises(BaseError):
            convert_files([path], output)
        self.assertEqual(len(list(self.folder.glob("*_DB_delete*.xlsx"))), 1)

    def test_repeated_names_allocate_numbers_without_overwriting(self):
        path = self.csv_file([record()], name="base.csv")
        first = convert_files([path], output_folder=self.folder)
        before = first.output.read_bytes()
        second = convert_files([path], output_folder=self.folder)
        self.assertEqual(first.output.name, "base_DB_delete_B1.xlsx")
        self.assertEqual(second.output.name, "base_DB_delete_B2.xlsx")
        self.assertEqual(first.output.read_bytes(), before)
        book = openpyxl.load_workbook(second.output, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Hoja1"]["K2"].value, "B2")
        finally:
            book.close()

    def test_known_number_keeps_identity_and_adds_version(self):
        path = self.csv_file([record()], name="base.csv")
        first = convert_files([path], base_number="b005", output_folder=self.folder)
        second = convert_files([path], base_number="5", output_folder=self.folder)
        self.assertEqual(first.output.name, "base_DB_delete_B5.xlsx")
        self.assertEqual(second.output.name, "base_DB_delete_B5_2.xlsx")

    def test_invalid_numbers_do_not_create_outputs(self):
        path = self.csv_file([record()], name="base.csv")
        for number in ["B0", "-1", "abc", "B1.5", "1234567890"]:
            with self.subTest(number=number), self.assertRaises(BaseError):
                convert_files([path], base_number=number, output_folder=self.folder)
        self.assertEqual(list(self.folder.glob("*.xlsx")), [])

    def test_origin_number_is_used_for_filename(self):
        path = self.csv_file([record() + ["B7"]], RAW_HEADERS + ["N° BASE"], name="base.csv")
        result = convert_files([path], output_folder=self.folder)
        self.assertEqual(result.output.name, "base_DB_delete_B7.xlsx")

    def test_collision_during_save_retries_and_updates_base_number(self):
        path = self.csv_file([record()], name="base.csv")
        original_writer = write_workbook
        calls = []
        def racing_writer(output, *args, **kwargs):
            calls.append(output)
            if len(calls) == 1:
                output.write_bytes(b"resultado de otro proceso")
            return original_writer(output, *args, **kwargs)
        with patch("scripts.transformar_base.write_workbook", side_effect=racing_writer):
            result = convert_files([path], output_folder=self.folder)
        self.assertEqual(calls[0].read_bytes(), b"resultado de otro proceso")
        self.assertEqual(result.output.name, "base_DB_delete_B2.xlsx")
        book = openpyxl.load_workbook(result.output, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Hoja1"]["K2"].value, "B2")
        finally:
            book.close()

    @unittest.skipUnless(os.name == "nt", "Prueba del guardado atómico de Windows")
    def test_destination_created_at_final_commit_is_not_overwritten(self):
        path = self.csv_file([record()])
        source = read_csv(path, None, None)[0]
        output = self.folder / "resultado.xlsx"
        rename = Path.rename
        def raced_rename(temporary, target):
            Path(target).write_bytes(b"archivo anterior")
            return rename(temporary, target)
        with patch.object(Path, "rename", raced_rename), self.assertRaises(OutputCollisionError):
            write_workbook(output, [source], [transform(source)])
        self.assertEqual(output.read_bytes(), b"archivo anterior")
        self.assertEqual(list(self.folder.glob("*.tmp.xlsx")), [])

    def test_unknown_base_does_not_take_known_number_in_consolidation(self):
        first = self.csv_file([record()], name="sin numero.csv")
        second = self.csv_file([record(phone="0990000002")], name="con numero B1.csv")
        result = convert_files([first, second], output_folder=self.folder)
        self.assertEqual(result.output.name, "sin numero_DB_delete_B2.xlsx")
        book = openpyxl.load_workbook(result.output, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Validar"]["K2"].value, "B2")
            self.assertEqual(book["Validar"]["K3"].value, "B1")
        finally:
            book.close()

    def test_model_ruc_exclusion_and_override(self):
        source = read_csv(self.csv_file([
            record(identifier="1791234567001"),
            record(phone="0990000002", identifier="0123456789001"),
        ]), None, None)[0]
        self.assertEqual(len(transform(source)), 1)
        self.assertEqual(len(transform(source, include_ruc_third_9=True)), 2)

    def test_names_collapse_padding_without_changing_original_data(self):
        source = read_csv(self.csv_file([record(name="PERSONA     DE   PRUEBA")]), None, None)[0]
        self.assertEqual(transform(source)[0][5], "PERSONA DE PRUEBA")
        self.assertEqual(source.rows[0][8], "PERSONA     DE   PRUEBA")

    def test_duplicated_surname_blocks_are_cleaned_but_equal_surnames_remain(self):
        source = read_csv(self.csv_file([
            record(name="PERSONA DE PRUEBA APELLIDO PRUEBA APELLIDO"),
            record(phone="0990000002", name="PERSONA DE PRUEBA PRUEBA"),
        ]), None, None)[0]
        self.assertEqual(transform(source)[0][5], "PERSONA DE PRUEBA APELLIDO")
        self.assertEqual(transform(source)[1][5], "PERSONA DE PRUEBA PRUEBA")

    def test_names_remove_standalone_period_but_preserve_initials(self):
        source = read_csv(self.csv_file([
            record(name="PERSONA DE PRUEBA ."),
            record(phone="0990000002", name="PERSONA DE PRUEBA A."),
        ]), None, None)[0]
        self.assertEqual(transform(source)[0][5], "PERSONA DE PRUEBA")
        self.assertEqual(transform(source)[1][5], "PERSONA DE PRUEBA A.")

    def test_excel_table_filters_colors_and_model_sheet_names(self):
        source = read_csv(self.csv_file([record()], name="campana B3.csv"), None, None)[0]
        output = self.folder / "estructura.xlsx"
        write_workbook(output, [source], [transform(source, base_type="R")])
        book = openpyxl.load_workbook(output, data_only=False)
        try:
            original = book["campana B3"]
            self.assertEqual(original["A1"].value, "Column1")
            self.assertEqual(original["K2"].value, "FORMULARIO")
            self.assertEqual(original["A3"].value, "Teléfono")
            self.assertEqual(original.tables["BaseOriginal1"].tableStyleInfo.name, "TableStyleMedium7")
            self.assertEqual(original.tables["BaseOriginal1"].ref, "A1:K4")
            self.assertEqual(book["Hoja3"].auto_filter.ref, "A1:K2")
            self.assertEqual(book["Validar"].auto_filter.ref, "A1:K2")
            self.assertEqual(len(book["Hoja3"].conditional_formatting), 2)
            self.assertIn("Hoja3", book.sheetnames)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
