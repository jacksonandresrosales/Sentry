"""Integración de bases Lucid con historial, búsqueda de audios y exportación."""
import csv
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from app.database import Database
from app.services.base_conversion import (
    BaseAudioIndex, load_hoja1_audio_index, normalize_phone_number, process_bases,
)
from app.services.report_export import export_audit_excel, load_base_export_records
from scripts.transformar_base import DEFAULT_STATE, HEADERS


class LucidServicesTests(unittest.TestCase):
    def test_phone_normalization_matches_mobile_and_fixed_line_variants(self):
        for national in ("990000001", "22345678", "59345678"):
            for value in (national, "0" + national, "593" + national,
                          "+593 " + national, "00593" + national):
                with self.subTest(value=value):
                    normalized = normalize_phone_number(value)
                    self.assertEqual(normalized, "0" + national)
                    self.assertEqual(normalize_phone_number(normalized), normalized)
        self.assertEqual(normalize_phone_number("(02) 234-5678"), "022345678")
        self.assertEqual(normalize_phone_number("+593 (99) 000-0001"), "0990000001")

    def test_phone_normalization_preserves_permissive_legacy_identifiers(self):
        for value, expected in ((None, ""), ("", ""), ("12345", "12345"),
                                ("ext. 12345", "12345"), ("+44 12345678901", "4412345678901")):
            with self.subTest(value=value):
                self.assertEqual(normalize_phone_number(value), expected)

    def make_workbook(self, path, headers, rows):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Hoja1"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()

    def make_lucid_base(self, path):
        self.make_workbook(path, ["Teléfono", "Nombre", "Estado"], [
            ["0990000001", "CLIENTE AUTOMÁTICO", "No contesta"],
            ["0980000002", "CLIENTE VERIFICADO", "Contactado"],
            ["0970000003", "CLIENTE NORMAL", "Reagendar"],
            ["+593 990 000 001", "NOMBRE REPETIDO", "ESTADO REPETIDO"],
        ])

    def test_lucid_conversion_persists_correct_columns_and_system(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "consolidado_2026-09-16.csv"
            with source.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["TELEFONO", "NOMBRE", "GESTION"])
                writer.writerow(["990000001", "CLIENTE UNO", "ELIMINAR DE BASE DE DATOS"])
                writer.writerow(["980000002", "CLIENTE DOS", "Reagendar"])
                writer.writerow(["0990000001", "DUPLICADO", "Contactado"])
            database = Database(folder / "audit.db")
            result = process_bases(database, [source], folder, source_system="lucid")
            self.assertEqual((result.filtered, result.unique, result.source_system), (3, 2, "lucid"))
            index = load_hoja1_audio_index(result.output)
            self.assertEqual(index.source_system, "lucid")
            self.assertEqual(index.names, {"0990000001": "CLIENTE UNO", "0980000002": "CLIENTE DOS"})
            self.assertEqual(index.phone_dates, {"0990000001": frozenset(), "0980000002": frozenset()})
            self.assertEqual(index.dates, set())
            job = database.jobs()[0]
            self.assertEqual(job["status"], "COMPLETADO")
            self.assertEqual(json.loads(job["options_json"])["source_system"], "lucid")
            with database.connect() as connection:
                rows = connection.execute("SELECT * FROM base_records ORDER BY row_number").fetchall()
                self.assertEqual(len(rows), 2)
                self.assertEqual((rows[0]["phone"], rows[0]["name"], rows[0]["state"]),
                                 ("0990000001", "CLIENTE UNO", "ELIMINAR DE BASE DE DATOS"))
                self.assertEqual(rows[1]["state"], "Reagendar")
                for row in rows:
                    self.assertEqual(row["base_type"], "V")
                    for key in ("call_state", "agent", "called_at", "duration", "identifier", "identifier_type", "base_number"):
                        self.assertEqual(row[key], "", key)
                self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")

    def test_lucid_index_reads_headers_not_positions_and_does_not_invent_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "CAMP_2026-09-16.xlsx"
            self.make_workbook(path, ["Estado", "Nombre", "Teléfono"], [
                ["Contactado", "PRIMER CLIENTE", "0990000001"],
                ["Otro estado", "SEGUNDO NOMBRE", "+593 990 000 001"],
                ["No contesta", "OTRO CLIENTE", "0980000002"],
            ])
            index = load_hoja1_audio_index(path)
            self.assertEqual(index.source_system, "lucid")
            self.assertEqual(index.phones, {"0990000001", "0980000002"})
            self.assertEqual(index.phone_dates, {"0990000001": frozenset(), "0980000002": frozenset()})
            self.assertEqual(index.names["0990000001"], "PRIMER CLIENTE")
            records = load_base_export_records(path)
            self.assertEqual(len(records), 2)
            self.assertEqual((records[0].customer, records[0].state, records[0].identifier),
                             ("PRIMER CLIENTE", "Contactado", ""))

    def test_issabel_reordered_headers_retain_names_and_actual_call_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "base.xlsx"
            row = ["0990000001", "Success", "001", "2026-09-15 10:30:00", "40",
                   "CLIENTE ISSABEL", "0123456789", "CÉDULA", DEFAULT_STATE, "V", "B1"]
            self.make_workbook(path, list(reversed(HEADERS)), [list(reversed(row))])
            index = load_hoja1_audio_index(path)
            self.assertEqual(index.source_system, "issabel")
            self.assertEqual(index.phone_dates, {"0990000001": frozenset({"20260915"})})
            self.assertEqual(index.names, {"0990000001": "CLIENTE ISSABEL"})
            record = load_base_export_records(path)[0]
            self.assertEqual((record.customer, record.identifier, record.source_system),
                             ("CLIENTE ISSABEL", "0123456789", "issabel"))
            self.assertEqual(BaseAudioIndex({}).source_system, "issabel")

    def test_incomplete_or_ambiguous_bases_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "base.xlsx"
            for headers in (["Teléfono", "Nombre"], ["Teléfono", "Nombre", "Estado", "Fecha"],
                            ["Teléfono", "Nombre", "Estado", "Teléfono"]):
                with self.subTest(headers=headers):
                    self.make_workbook(path, headers, [])
                    with self.assertRaisesRegex(ValueError, "Hoja1"):
                        load_hoja1_audio_index(path)
                    with self.assertRaisesRegex(ValueError, "Hoja1"):
                        load_base_export_records(path)

    def test_old_lucid_export_modes_add_empty_id_and_preserve_phone_text(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            base = folder / "base.xlsx"
            self.make_lucid_base(base)
            automatic = {"990000001", "+593 980 000 002"}
            verified = {"980000002"}
            expected = {
                "automatic": [("0990000001", "CLIENTE AUTOMÁTICO", "", "Denuncia detectada")],
                "verified": [("0980000002", "CLIENTE VERIFICADO", "", "Denuncia verificada")],
                "complaints": [("0990000001", "CLIENTE AUTOMÁTICO", "", "Denuncia detectada"),
                               ("0980000002", "CLIENTE VERIFICADO", "", "Denuncia verificada")],
                "all": [("0990000001", "CLIENTE AUTOMÁTICO", "", "Denuncia detectada"),
                        ("0980000002", "CLIENTE VERIFICADO", "", "Denuncia verificada"),
                        ("0970000003", "CLIENTE NORMAL", "", "Reagendar")],
            }
            for mode, rows in expected.items():
                with self.subTest(mode=mode):
                    result = export_audit_excel(folder / f"{mode}.xlsx", base, automatic, verified, mode)
                    self.assertEqual(result.row_count, len(rows))
                    self.assertEqual(result.complaint_count, min(2, len(rows)))
                    book = openpyxl.load_workbook(result.output, data_only=False)
                    try:
                        sheet = book["Resultados"]
                        self.assertEqual(tuple(cell.value for cell in sheet[1]), ("Teléfono", "Nombre", "ID", "Estado"))
                        self.assertEqual(list(sheet.iter_rows(min_row=2, values_only=True)), rows)
                        for cell in sheet["A"][1:]:
                            self.assertEqual(cell.data_type, "s")
                            self.assertEqual(cell.number_format, "@")
                        self.assertEqual(sheet.max_column, 4)
                    finally:
                        book.close()

    def test_issabel_export_modes_keep_legacy_columns_and_state(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            base = folder / "base.xlsx"
            self.make_workbook(base, HEADERS, [
                [f"09{number}000000{number}", "Success", "001", "2026-09-15", "30",
                 f"CLIENTE {number}", f"012345678{number}", "CÉDULA", DEFAULT_STATE, "V", "B1"]
                for number in (1, 2, 3)
            ])
            for mode, expected in (("automatic", [1]), ("verified", [2]), ("complaints", [1, 2]), ("all", [1, 2, 3])):
                with self.subTest(mode=mode):
                    result = export_audit_excel(folder / f"{mode}.xlsx", base,
                                                {"0910000001", "0920000002"}, {"0920000002"}, mode)
                    book = openpyxl.load_workbook(result.output, data_only=True)
                    try:
                        sheet = book["Resultados"]
                        self.assertEqual(tuple(cell.value for cell in sheet[1]),
                                         ("Número de celular", "Nombre del cliente", "ID", "Estado"))
                        self.assertEqual(list(sheet.iter_rows(min_row=2, values_only=True)), [
                            (f"09{number}000000{number}", f"CLIENTE {number}", f"012345678{number}", DEFAULT_STATE)
                            for number in expected
                        ])
                    finally:
                        book.close()

    def test_export_cannot_replace_active_base_even_after_extension_correction(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp) / "base.xlsx"
            self.make_lucid_base(base)
            original = base.read_bytes()
            for output in (base, base.with_suffix(".csv")):
                with self.subTest(output=output):
                    with self.assertRaisesRegex(ValueError, "reemplazar la base activa"):
                        export_audit_excel(output, base, {"0990000001"}, set(), "all")
                    self.assertEqual(base.read_bytes(), original)

    def test_new_lucid_ids_survive_conversion_database_export_and_header_reordering(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "lucid.csv"
            with source.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["ID", "Estado", "Celular", "Nombre"])
                writer.writerow(["0012345678", "Contactado", "990000001", "PERSONA UNO"])
                writer.writerow(["0012345678001", "Pendiente", "990000002", "PERSONA CON RUC"])
                writer.writerow(["1790000000001", "Pendiente", "990000003", "EMPRESA S.A."])
            original = source.read_bytes()
            database = Database(folder / "audit.db")
            result = process_bases(database, [source], folder, source_system="lucid")
            self.assertEqual((result.filtered, result.unique, result.companies, result.excluded), (2, 2, 1, 1))
            self.assertEqual(json.loads(database.jobs()[0]["options_json"])["entity_filter"], "people")
            with database.connect() as connection:
                rows = connection.execute("SELECT identifier,identifier_type FROM base_records ORDER BY row_number").fetchall()
                self.assertEqual([tuple(row) for row in rows], [("0012345678", "CÉDULA"), ("0012345678001", "RUC")])
            records = load_base_export_records(result.output)
            self.assertEqual([record.identifier for record in records], ["0012345678", "0012345678001"])
            exported = export_audit_excel(folder / "export.xlsx", result.output, {"990000001"}, set(), "all")
            book = openpyxl.load_workbook(exported.output, data_only=False)
            try:
                sheet = book["Resultados"]
                self.assertEqual(list(sheet.iter_rows(values_only=True)), [
                    ("Teléfono", "Nombre", "ID", "Estado"),
                    ("0990000001", "PERSONA UNO", "0012345678", "Denuncia detectada"),
                    ("0990000002", "PERSONA CON RUC", "0012345678001", "Pendiente"),
                ])
                self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in sheet["C"][1:]))
            finally:
                book.close()
            reordered = folder / "reordered.xlsx"
            self.make_workbook(reordered, ["ID", "Estado", "Teléfono", "Nombre"], [
                ["0012345678", "Pendiente", "0990000001", "PERSONA UNO"],
            ])
            self.assertEqual(load_base_export_records(reordered)[0].identifier, "0012345678")
            self.assertEqual(load_hoja1_audio_index(reordered).phones, {"0990000001"})
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
