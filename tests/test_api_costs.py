"""Facturación local: una solicitud real se cuenta una vez aun con caché."""
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import openpyxl

from app.database import Database
from app.services import audio_analysis
from app.services.api_cost_export import export_api_costs_excel
from app.services.api_costs import load_cost_report


class ApiCostsTests(unittest.TestCase):
    def test_automatic_tariffs_and_excel_totals(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            database.record_api_usage(path, 'transcription', 'deepgram', 'nova-3',
                                      {'metadata': {'duration': 60, 'request_id': 'dg-1'},
                                       'usage': {'input_tokens': 40, 'output_tokens': 10}}, 60,
                                      keyterm=True)
            database.record_api_usage(path, 'analysis', 'gemini', 'gemini-3.5-flash',
                                      {'usageMetadata': {'promptTokenCount': 1000,
                                                         'candidatesTokenCount': 100}}, None)
            for hour in (2, 12):
                created = datetime(2026, 9, 21, hour, tzinfo=timezone.utc)
                database.record_api_usage(path, 'analysis', 'deepseek', 'deepseek-flash',
                                          {'created': int(created.timestamp()),
                                           'usage': {'prompt_tokens': 1000, 'completion_tokens': 100}}, None)
            with database.connect() as connection:
                connection.execute("UPDATE api_usage SET created_at='2026-09-23 12:00:00'")
            start, end = date(2026, 9, 23), date(2026, 9, 24)
            report = load_cost_report(database, start, end)
            self.assertAlmostEqual(report['total_usd'], 0.00863)
            output = export_api_costs_excel(database, Path(folder) / 'gastos.xlsx', start, end)
            book = openpyxl.load_workbook(output, data_only=True)
            try:
                self.assertEqual(book.sheetnames, ['Resumen', 'Diario', 'Solicitudes', 'Tarifas'])
                self.assertAlmostEqual(book['Resumen']['B12'].value, report['total_usd'])
                self.assertAlmostEqual(book['Solicitudes']['E6'].value, report['total_usd'])
                self.assertAlmostEqual(book['Tarifas']['I22'].value, report['total_usd'])
                self.assertEqual(book['Solicitudes'].max_row, 6)
                self.assertIsInstance(book['Solicitudes']['A2'].value, datetime)
                self.assertIsInstance(book['Diario']['A2'].value, datetime)
                self.assertEqual(book['Solicitudes'].freeze_panes, 'E2')
                self.assertEqual(book['Solicitudes']['F1'].value, 'Estado')
                self.assertEqual(book['Solicitudes'].auto_filter.ref, 'A1:Q5')
                self.assertTrue(book['Solicitudes'].column_dimensions['L'].hidden)
                self.assertEqual(len(book['Resumen']._charts), 1)
                chart = book['Resumen']._charts[0]
                self.assertEqual(chart.series[0].marker.symbol, 'circle')
                self.assertEqual([series.val.numRef.numCache.ptCount for series in chart.series], [1, 1])
                self.assertAlmostEqual(sum(series.val.numRef.numCache.pt[0].v for series in chart.series),
                                       report['total_usd'])
                self.assertAlmostEqual(sum(book['Resumen'].cell(row, 7).value for row in range(18, 21)), 1)
                self.assertAlmostEqual(book['Resumen']['H18'].value, 1)
                self.assertEqual(book['Resumen']['I18'].value, 40)
                self.assertEqual(book['Resumen']['J18'].value, 10)
                self.assertFalse(book['Resumen'].sheet_view.showGridLines)
            finally:
                book.close()

    def test_excel_preserves_missing_costs_zero_costs_and_literal_identifiers(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / '=SUM(1,1).wav'
            with database.connect() as connection:
                connection.execute('INSERT INTO calls(filename,file_path) VALUES (?,?)', (path.name, str(path.resolve())))
            database.record_api_usage(path, 'analysis', 'gemini', '=1+1',
                                      {'usageMetadata': {'promptTokenCount': 500, 'candidatesTokenCount': 30}}, None)
            database.record_api_usage(path, 'analysis', 'groq', 'openai/gpt-oss-120b',
                                      {'usage': {'prompt_tokens': 0, 'completion_tokens': 0}}, None)
            today = date.today()
            output = export_api_costs_excel(database, Path(folder) / 'costs.xlsx', today, today + timedelta(days=1))
            book = openpyxl.load_workbook(output, data_only=True)
            try:
                self.assertEqual(book['Solicitudes']['E2'].value, 0)
                self.assertEqual(book['Solicitudes']['F2'].value, 'Estimado')
                self.assertIsNone(book['Solicitudes']['E3'].value)
                self.assertEqual(book['Solicitudes']['F3'].value, 'Sin estimación')
                self.assertEqual(book['Solicitudes']['D3'].value, '=1+1')
                self.assertEqual(book['Solicitudes']['D3'].data_type, 's')
                self.assertEqual(book['Solicitudes']['G3'].value, path.name)
                self.assertEqual(book['Solicitudes']['G3'].data_type, 's')
                self.assertEqual(book['Resumen']['B9'].value, 1)
                self.assertEqual(book['Resumen']['B12'].value, 0)
            finally:
                book.close()

    def test_excel_exports_all_filtered_requests_and_reconciles_totals(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            with database.connect() as connection:
                connection.executemany("""INSERT INTO api_usage
                    (service,provider,model,request_id,cost_usd,created_at)
                    VALUES ('analysis','groq','openai/gpt-oss-120b',?,?,datetime(?,'utc'))""",
                    [(str(index), .001, '2026-10-01 12:00:00') for index in range(502)] +
                    [('outside-period', 100, '2026-10-02 12:00:00')])
                connection.execute("""INSERT INTO api_usage (service,provider,model,cost_usd,created_at)
                    VALUES ('analysis','gemini','custom',200,datetime('2026-10-01 12:00:00','utc'))""")
            output = export_api_costs_excel(database, Path(folder) / 'costs.xlsx', date(2026, 10, 1),
                                           date(2026, 10, 2), 'analysis', 'groq')
            book = openpyxl.load_workbook(output, data_only=True)
            try:
                self.assertEqual(book['Solicitudes'].max_row, 504)
                self.assertEqual(book['Resumen']['B8'].value, 502)
                for sheet, cell in (('Resumen', 'B12'), ('Resumen', 'E19'), ('Diario', 'D3'),
                                    ('Solicitudes', 'E504'), ('Tarifas', 'I19')):
                    self.assertAlmostEqual(book[sheet][cell].value, .502, msg=f'{sheet}!{cell}')
                self.assertEqual(book['Resumen']['B5'].value, 'Análisis')
                self.assertEqual(book['Resumen']['B6'].value, 'Groq')
            finally:
                book.close()

    def test_empty_excel_has_zero_totals_without_self_referencing_formulas(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            day = date(2026, 10, 1)
            output = export_api_costs_excel(database, Path(folder) / 'empty.xlsx', day, day)
            book = openpyxl.load_workbook(output)
            try:
                self.assertEqual(book['Diario']['D2'].value, 0)
                self.assertEqual(book['Solicitudes']['E2'].value, 0)
                self.assertEqual(book['Resumen']['E18'].value, 0)
                self.assertEqual(book['Tarifas']['I18'].value, 0)
                self.assertEqual(len(book['Resumen']._charts), 0)
            finally:
                book.close()

    def test_assemblyai_rates_include_diarization_and_keyterms(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            database.record_api_usage(path, 'transcription', 'assemblyai', 'universal-2',
                                      {'id': 'aai-1', 'audio_duration': 3600}, None, keyterm=True)
            database.record_api_usage(path, 'transcription', 'assemblyai', 'universal-3-5-pro',
                                      {'id': 'aai-2', 'audio_duration': 3600}, None, keyterm=True)
            today = date.today()
            report = load_cost_report(database, today, today + timedelta(days=1))
            costs = {event['model']: event['cost_usd'] for event in report['events']}
            self.assertAlmostEqual(costs['universal-2'], 0.17)
            self.assertAlmostEqual(costs['universal-3-5-pro'], 0.28)
            self.assertEqual({event['request_id'] for event in report['events']}, {'aai-1', 'aai-2'})

    def test_assemblyai_usage_is_recorded_once_when_transcript_completes(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            path.write_bytes(b'RIFF')
            responses = [{'upload_url': 'https://cdn.assemblyai.com/upload/x'}, {'id': 'tr-1', 'status': 'queued'},
                         {'id': 'tr-1', 'status': 'completed', 'text': 'hola', 'audio_duration': 120}]
            context = audio_analysis._COST_CONTEXT.set({
                'database': database, 'path': path, 'duration_seconds': 120,
                'config': {'transcription_model': 'universal-2'},
            })
            try:
                with patch('app.services.audio_analysis._request', side_effect=responses), \
                        patch('app.services.audio_analysis.time.sleep'):
                    audio_analysis._assemblyai_transcript(path, 'k', 'universal-2', ['demanda'])
            finally:
                audio_analysis._COST_CONTEXT.reset(context)
            today = date.today()
            report = load_cost_report(database, today, today + timedelta(days=1))
            self.assertEqual(len(report['events']), 1)
            self.assertAlmostEqual(report['events'][0]['cost_usd'], 2 * 0.17 / 60)

    def test_assemblyai_charges_the_model_returned_by_the_provider(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            context = audio_analysis._COST_CONTEXT.set({
                'database': database, 'path': Path(folder) / 'audio.wav', 'duration_seconds': 3600,
                'config': {'transcription_model': 'universal-3-5-pro'},
            })
            try:
                audio_analysis._record_usage('transcription', 'assemblyai',
                    {'speech_model_used': 'universal-2', 'audio_duration': 3600}, keyterm=True)
            finally:
                audio_analysis._COST_CONTEXT.reset(context)
            today = date.today()
            event = load_cost_report(database, today, today + timedelta(days=1))['events'][0]
            self.assertEqual(event['model'], 'universal-2')
            self.assertAlmostEqual(event['cost_usd'], 0.17)

    def test_request_cache_and_historical_rate(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            with database.connect() as connection:
                connection.execute('INSERT INTO calls(filename,file_path,duration_seconds) VALUES (?,?,?)',
                                   (path.name, str(path), 60))
            response = SimpleNamespace(status_code=200, headers={}, json=lambda: {
                'metadata': {'duration': 60, 'request_id': 'provider-request-1'}, 'results': {}})
            session = Mock(request=Mock(return_value=response))
            context = audio_analysis._COST_CONTEXT.set({
                'database': database, 'path': path, 'duration_seconds': 60,
                'config': {'transcription_model': 'nova-3'},
            })
            try:
                with patch.object(audio_analysis._HTTP, 'session', session, create=True):
                    compute = lambda: audio_analysis._request('POST', 'https://api.deepgram.com/v1/listen')
                    first, cached_first = audio_analysis._cached_computation(database, 'transcription_cache', 'audio-1', compute)
                    second, cached_second = audio_analysis._cached_computation(database, 'transcription_cache', 'audio-1', compute)
            finally:
                audio_analysis._COST_CONTEXT.reset(context)
            self.assertEqual(first, second)
            self.assertFalse(cached_first)
            self.assertTrue(cached_second)
            session.request.assert_called_once()
            today = date.today()
            report = load_cost_report(database, today, today + timedelta(days=1))
            self.assertEqual(report['requests'], 1)
            self.assertAlmostEqual(report['total_usd'], .0043)
            database.save_api_price('transcription', 'deepgram', 'nova-3', 'minute', .01, 0)
            self.assertAlmostEqual(load_cost_report(database, today, today + timedelta(days=1))['total_usd'], .0043)
            database.record_api_usage(path, 'transcription', 'deepgram', 'nova-3',
                                      {'metadata': {'duration': 60, 'request_id': 'provider-request-2'}}, 60)
            self.assertAlmostEqual(load_cost_report(database, today, today + timedelta(days=1))['total_usd'], .0086)

    def test_missing_rate_keeps_usage_visible_without_inventing_cost(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            database.record_api_usage(Path(folder) / 'audio.wav', 'analysis', 'gemini', 'custom-model',
                                      {'usageMetadata': {'promptTokenCount': 500, 'candidatesTokenCount': 30}}, None)
            today = date.today()
            report = load_cost_report(database, today, today + timedelta(days=1))
            self.assertEqual(report['requests'], 1)
            self.assertEqual(report['unpriced'], 1)
            self.assertIsNone(report['events'][0]['cost_usd'])

    def test_daily_chart_breaks_down_both_services_and_models(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            database.record_api_usage(path, 'transcription', 'deepgram', 'nova-3',
                                      {'metadata': {'duration': 60}}, 60)
            database.record_api_usage(path, 'analysis', 'gemini', 'gemini-3.5-flash-lite',
                                      {'usageMetadata': {'promptTokenCount': 1000,
                                                         'candidatesTokenCount': 100}}, None)
            with database.connect() as connection:
                connection.execute("UPDATE api_usage SET created_at=datetime('2026-09-21 12:00:00','utc')")
            report = load_cost_report(database, date(2026, 9, 21), date(2026, 9, 24))
            self.assertEqual(len(report['daily']), 3)
            day = report['daily'][0]
            self.assertEqual(day['date'], '2026-09-21')
            self.assertEqual(day['transcription_requests'], 1)
            self.assertEqual(day['analysis_requests'], 1)
            self.assertAlmostEqual(day['transcription_usd'], .0043)
            self.assertAlmostEqual(day['total_usd'],
                                   day['transcription_usd'] + day['analysis_usd'])
            self.assertEqual({row['model'] for row in day['models']},
                             {'nova-3', 'gemini-3.5-flash-lite'})
            self.assertEqual(report['daily'][1]['total_usd'], 0)

    def test_report_aggregates_consumption_per_model_and_limits_listed_events(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            path = Path(folder) / 'audio.wav'
            with database.connect() as connection:
                connection.execute('INSERT INTO calls(filename,file_path) VALUES (?,?)', (path.name, str(path.resolve())))
            for _ in range(3):
                database.record_api_usage(path, 'transcription', 'assemblyai', 'universal-2',
                                          {'audio_duration': 1800}, None)
            database.record_api_usage(path, 'analysis', 'groq', 'openai/gpt-oss-120b',
                                      {'usage': {'prompt_tokens': 2000, 'completion_tokens': 500}}, None)
            today = date.today()
            report = load_cost_report(database, today, today + timedelta(days=1), event_limit=2)
            self.assertEqual(len(report['events']), 2)
            self.assertEqual(report['requests'], 4)
            self.assertAlmostEqual(report['audio_seconds'], 5400)
            universal = next(item for item in report['by_model'] if item['model'] == 'universal-2')
            self.assertEqual(universal['label'], 'AssemblyAI Universal-2')
            self.assertAlmostEqual(universal['usd_per_hour'], 0.17)
            self.assertAlmostEqual(sum(item['share'] for item in report['by_model']), 1)
            groq = next(item for item in report['by_model'] if item['provider'] == 'groq')
            self.assertEqual((groq['input_tokens'], groq['output_tokens']), (2000, 500))
            self.assertIsNone(groq['usd_per_hour'])
            self.assertEqual(report['calls'], 1)

    def test_long_periods_are_grouped_by_week_or_month(self):
        from app.services.api_costs import bucket_daily
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            database.record_api_usage(Path(folder) / 'a.wav', 'transcription', 'deepgram', 'nova-3',
                                      {'metadata': {'duration': 60}}, 60)
            with database.connect() as connection:
                connection.execute("UPDATE api_usage SET created_at=datetime('2026-03-04 12:00:00','utc')")
            report = load_cost_report(database, date(2026, 1, 1), date(2026, 7, 1))
            granularity, rows = bucket_daily(report['daily'])
            self.assertEqual(granularity, 'week')
            self.assertAlmostEqual(sum(row['total_usd'] for row in rows), report['total_usd'])
            week = next(row for row in rows if row['total_usd'])
            self.assertEqual((week['date'], week['end']), ('2026-03-02', '2026-03-08'))
            self.assertEqual(week['models'][0]['model'], 'nova-3')
            granularity, rows = bucket_daily(report['daily'], 'month')
            self.assertEqual([row['label'] for row in rows][:3], ['01/2026', '02/2026', '03/2026'])

    def test_effective_hourly_cost_excludes_audio_without_a_cost_estimate(self):
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            for index in range(2):
                database.record_api_usage(Path(folder) / f'{index}.wav', 'transcription', 'assemblyai',
                                          'universal-2', {'audio_duration': 3600}, None)
            with database.connect() as connection:
                connection.execute('UPDATE api_usage SET cost_usd=NULL WHERE id=(SELECT MAX(id) FROM api_usage)')
            today = date.today()
            model = load_cost_report(database, today, today + timedelta(days=1))['by_model'][0]
            self.assertEqual(model['audio_seconds'], 7200)
            self.assertEqual(model['unpriced'], 1)
            self.assertAlmostEqual(model['usd_per_hour'], 0.17)


class ApiCostsPageTests(unittest.TestCase):
    def test_empty_page_and_saved_usage_match_the_database(self):
        import os
        os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
        from PySide6.QtCore import QEvent
        from PySide6.QtWidgets import QApplication
        from app.ui.views.api_costs_page import ApiCostsPage
        app = QApplication.instance() or QApplication([])
        with TemporaryDirectory() as folder:
            database = Database(Path(folder) / 'costs.db')
            page = ApiCostsPage(database)
            today = date.today()
            page.render(load_cost_report(database, today, today + timedelta(days=1)))
            self.assertEqual(page.metrics['total_usd'].text(), '$0.00')
            self.assertEqual(page.events.rowCount(), 0)
            self.assertIn('Aún no hay solicitudes', page.summary.text())
            path = Path(folder) / 'audio.wav'
            database.register_calls([SimpleNamespace(filename=path.name, source_path=path, duration=60)])
            database.record_api_usage(path, 'transcription', 'deepgram', 'nova-3',
                                      {'metadata': {'duration': 60}}, 60)
            page.render(load_cost_report(database, today, today + timedelta(days=1)))
            self.assertEqual(page.metrics['total_usd'].text(), '$0.0043')
            self.assertEqual(page.events.rowCount(), 1)
            self.assertEqual(page.calls.rowCount(), 1)
            self.assertEqual(page.events.item(0, 3).text(), 'audio.wav')
            self.assertEqual(page.provider.accessibleName(), 'Proveedor')
            page.close()
            page.deleteLater()
            app.sendPostedEvents(None, QEvent.Type.DeferredDelete)


if __name__ == '__main__':
    unittest.main()
