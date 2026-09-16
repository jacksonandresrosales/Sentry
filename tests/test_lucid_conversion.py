"""Conversión Lucid con datos ficticios y compatibilidad con el formato Issabel."""

import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from scripts.transformar_base import (
    BaseError,
    LUCID_HEADERS,
    convert_files,
    default_output,
    main,
    is_company_name,
    normalize_lucid_phone,
    read_csv,
    write_lucid_workbook,
)


class LucidConversionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def csv_file(self, rows, headers=None, filename="campana.csv"):
        path = self.folder / filename
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers or ["Celular", "Nombre", "GESTION"])
            writer.writerows(rows)
        return path

    def read_output(self, path):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, ["Hoja1"])
            sheet = workbook["Hoja1"]
            self.assertEqual(sheet.max_column, 4)
            self.assertEqual(tuple(cell.value for cell in sheet[1]), LUCID_HEADERS)
            rows = list(sheet.iter_rows(min_row=2))
            for row in rows:
                self.assertEqual(row[0].data_type, "s")
            return [tuple(cell.value for cell in row) for row in rows]
        finally:
            workbook.close()

    def test_only_public_columns_all_states_and_first_phone_are_preserved(self):
        source = self.csv_file([
            ["990000001", " PERSONA   DE PRUEBA . ", "VENTA", "1791234567001", "privado"],
            ["0990000001", "SEGUNDA PERSONA", "NO CONTESTA", "id2", "secreto"],
            ["+593 99-000-0002", "PERSONA DOS", "BUZON", "id3", "secreto"],
            ["0990000003", "PERSONA TRES", "DENUNCIA", "id4", "secreto"],
            ["0990000004", "PERSONA CUATRO", "ELIMINAR DE BASE DE DATOS", "id5", "secreto"],
        ], headers=["Celular", "Nombre", "GESTION", "ID", "NOTAS"])
        before = source.read_bytes()
        result = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual((result.filtered, result.unique, result.source_system), (5, 4, "lucid"))
        self.assertFalse(result.missing_base_number)
        self.assertEqual(result.sources, [("campana", 5, 5)])
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PERSONA DE PRUEBA", "1791234567001", "VENTA"),
            ("0990000002", "PERSONA DOS", "id3", "BUZON"),
            ("0990000003", "PERSONA TRES", "id4", "DENUNCIA"),
            ("0990000004", "PERSONA CUATRO", "id5", "ELIMINAR DE BASE DE DATOS"),
        ])
        self.assertEqual(source.read_bytes(), before)

    def test_export_base_uses_estado_not_sub_estado_or_callback_date(self):
        headers = [
            "Celular", "Nombre", "ID", "Forma_de_pago", "Plan", "Plan_Anterior",
            "Mejora", "Plan_Actual", "ciclo_facturacion", "Cantidad_de_Bono", "Precio",
            "Estado", "Sub-Estado", "Notas", "Fecha Rellamada",
        ]
        rows = [
            ["990000001", " PERSONA   UNO ", "id1", "pago", "plan", "anterior", "mejora",
             "actual", "ciclo", "bono", "precio", "ELIMINAR DE LA BASE", "VENTA",
             "nota privada", "2026-09-17 09:00:00"],
            ["0990000002", "PERSONA DOS", "id2", "", "", "", "", "", "", "", "",
             "PENDIENTE", "ELIMINAR DE LA BASE", "otra nota", ""],
            ["0990000003", "PERSONA TRES", "id3", "", "", "", "", "", "", "", "",
             "", "BUZON", "", ""],
        ]
        source = self.csv_file(rows, headers=headers, filename="export_base_campana.csv")
        before = source.read_bytes()
        parsed = read_csv(source, None, None, source_system="lucid")[0]
        self.assertEqual(parsed.mapping, {"Teléfono": 0, "Nombre": 1, "ID": 2, "Estado": 11})

        result = convert_files([source], output_folder=self.folder, source_system="lucid")

        self.assertEqual((result.filtered, result.unique), (3, 3))
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PERSONA UNO", "id1", "ELIMINAR DE LA BASE"),
            ("0990000002", "PERSONA DOS", "id2", "PENDIENTE"),
            ("0990000003", "PERSONA TRES", "id3", ""),
        ])
        self.assertEqual(source.read_bytes(), before)

    def test_export_base_estado_is_detected_independently_of_column_order(self):
        source = self.csv_file(
            [["NO USAR", "PERSONA UNO", "ELIMINAR DE LA BASE", "990000001"]],
            headers=["Sub-Estado", "Nombre", "Estado", "Celular"],
            filename="export_base_orden_distinto.csv",
        )
        result = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PERSONA UNO", "", "ELIMINAR DE LA BASE"),
        ])

    def test_export_base_and_previous_consolidado_can_be_combined(self):
        exported = self.csv_file(
            [["990000001", "PRIMERO", "ELIMINAR DE LA BASE", "NO USAR"]],
            headers=["Celular", "Nombre", "Estado", "Sub-Estado"],
            filename="export_base.csv",
        )
        consolidated = self.csv_file(
            [["0990000001", "DUPLICADO", "BUZON"],
             ["990000002", "SEGUNDO", "NO CONTESTA"]],
            filename="consolidado.csv",
        )
        result = convert_files([exported, consolidated], output_folder=self.folder,
                               source_system="lucid")
        self.assertEqual((result.filtered, result.unique), (3, 2))
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PRIMERO", "", "ELIMINAR DE LA BASE"),
            ("0990000002", "SEGUNDO", "", "NO CONTESTA"),
        ])

    def test_phone_zero_and_ecuador_prefix_normalization_is_idempotent(self):
        for value in ["990000001", 990000001, "0990000001", "+593990000001",
                      "593990000001", "00593990000001", "+593 (099) 000-0001"]:
            with self.subTest(value=value):
                normalized = normalize_lucid_phone(value)
                self.assertEqual(normalized, "0990000001")
                self.assertEqual(normalize_lucid_phone(normalized), normalized)
        self.assertEqual(normalize_lucid_phone("22345678"), "022345678")
        self.assertEqual(normalize_lucid_phone("59345678"), "059345678")

    def test_invalid_phones_abort_before_creating_an_output(self):
        for value in ["", "sin telefono", "099ABC0001", "990000001.0", "9.9E+08",
                      "+1990000001", "000990000001", "123", "0000000000"]:
            with self.subTest(value=value):
                source = self.csv_file([[value, "PERSONA", "VENTA"]])
                before = source.read_bytes()
                with self.assertRaisesRegex(BaseError, "campana, registro 2"):
                    convert_files([source], output_folder=self.folder, source_system="lucid")
                self.assertEqual(source.read_bytes(), before)
                self.assertFalse(list(self.folder.glob("*.xlsx")))

    def test_xlsx_numeric_phone_and_existing_zero_format_are_supported(self):
        path = self.folder / "origen.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Teléfono", "Nombre", "Estado"])
        sheet.append([990000001, "PERSONA UNO", "NO CONTESTA"])
        sheet["A2"].number_format = "0000000000"
        sheet.append([990000002, "PERSONA DOS", "GESTION PENDIENTE"])
        private = workbook.create_sheet("Notas internas")
        private.append(["No publicar", "Contenido privado"])
        workbook.save(path)
        workbook.close()
        before = path.read_bytes()
        result = convert_files([path], output_folder=self.folder, source_system="lucid")
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PERSONA UNO", "", "NO CONTESTA"),
            ("0990000002", "PERSONA DOS", "", "GESTION PENDIENTE"),
        ])
        self.assertEqual(path.read_bytes(), before)

    def test_minimal_aliases_cli_and_formula_like_values_are_text(self):
        source = self.csv_file([["990000001", '=HYPERLINK("example")', "=1+1"]],
                               headers=["TELEFONO", "NOMBRE", "GESTIÓN"])
        output = self.folder / "salida.xlsx"
        self.assertEqual(main([str(source), "--sistema", "lucid", "-o", str(output)]), 0)
        self.assertEqual(self.read_output(output), [("0990000001", '=HYPERLINK("example")', "", "=1+1")])
        workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
        try:
            self.assertEqual(workbook["Hoja1"]["B2"].data_type, "s")
            self.assertEqual(workbook["Hoja1"]["D2"].data_type, "s")
        finally:
            workbook.close()

    def test_escaped_quotes_after_delimiter_sample_are_preserved(self):
        quoted_name = 'PERSONA "APODO" DE PRUEBA'
        for system in ("lucid", "issabel"):
            with self.subTest(system=system):
                if system == "lucid":
                    headers = ["Celular", "Nombre", "GESTION"]
                    ordinary = ["0990000001", "PERSONA " + "A" * 120, "VENTA"]
                else:
                    headers = ["Teléfono", "Nombre", "ESTADO", "Estado Llamada", "Agente",
                               "Fecha y hora", "Duración(Seg)", "ID"]
                    ordinary = ["0990000001", "PERSONA " + "A" * 120, "ESTADO", "Success",
                                "Agente", "2026-09-16 10:00:00", "20", "0123456789"]
                last = ordinary.copy()
                last[1] = quoted_name
                source = self.csv_file([ordinary] * 1500 + [last], headers=headers)
                before = source.read_bytes()
                self.assertGreater(before.index(b'""APODO""'), 131072)
                parsed = read_csv(source, None, None, source_system=system)[0]
                self.assertEqual(len(parsed.rows), 1501)
                self.assertEqual(parsed.rows[-1][parsed.mapping["Nombre"]], quoted_name)
                self.assertEqual(source.read_bytes(), before)

    def test_wrong_system_is_explained_and_issabel_remains_the_default(self):
        lucid = self.csv_file([["0990000001", "PERSONA", "VENTA"]])
        with self.assertRaisesRegex(BaseError, "Selecciona el sistema Lucid"):
            convert_files([lucid], output_folder=self.folder)
        exported = self.csv_file(
            [["0990000001", "PERSONA", "ELIMINAR DE LA BASE", ""]],
            headers=["Celular", "Nombre", "Estado", "Sub-Estado"],
            filename="export_base.csv",
        )
        with self.assertRaisesRegex(BaseError, "Selecciona el sistema Lucid"):
            convert_files([exported], output_folder=self.folder, source_system="issabel")
        issabel = self.csv_file([
            ["0990000001", "Success", "Agente", "2026-09-16 10:00:00", "20",
             "PERSONA", "0123456789", "ELIMINAR DE BASE DE DATOS"],
        ], headers=["Teléfono", "Estado Llamada", "Agente", "Fecha y hora", "Duración(Seg)",
                    "Nombre", "ID", "ESTADO"], filename="issabel.csv")
        with self.assertRaisesRegex(BaseError, "Selecciona el sistema Issabel"):
            convert_files([issabel], output_folder=self.folder, source_system="lucid")
        legacy = convert_files([issabel], output_folder=self.folder)
        self.assertEqual(legacy.source_system, "issabel")
        self.assertEqual((legacy.filtered, legacy.unique), (1, 1))

    def test_missing_or_duplicate_columns_are_reported(self):
        source = self.csv_file([["0990000001", "PERSONA"]], headers=["Celular", "Nombre"])
        with self.assertRaisesRegex(BaseError, "Columnas faltantes: Estado"):
            convert_files([source], output_folder=self.folder, source_system="lucid")
        source = self.csv_file([["0990000001", "PERSONA", "BUZON"]],
                               headers=["Celular", "Nombre", "Sub-Estado"])
        with self.assertRaisesRegex(BaseError, "Columnas faltantes: Estado"):
            convert_files([source], output_folder=self.folder, source_system="lucid")
        source = self.csv_file([["0990000001", "PERSONA", "VENTA", "BUZON"]],
                               headers=["Celular", "Nombre", "Estado", "GESTION"])
        with self.assertRaisesRegex(BaseError, "columna 'Estado' aparece más de una vez"):
            convert_files([source], output_folder=self.folder, source_system="lucid")
        source = self.csv_file([["0990000001", "0990000002", "PERSONA", "VENTA"]],
                               headers=["Celular", "Telefono", "Nombre", "GESTION"])
        with self.assertRaisesRegex(BaseError, "aparece más de una vez"):
            convert_files([source], output_folder=self.folder, source_system="lucid")

    def test_auto_names_preserve_previous_output_and_ignore_issabel_options(self):
        source = self.csv_file([["0990000001", "PERSONA", "BUZON"]])
        first = convert_files([source], output_folder=self.folder, source_system="lucid",
                              base_number="B0", base_type="invalid", state="inexistente",
                              call_state="Error")
        before = first.output.read_bytes()
        second = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual(first.output.name, "campana_Lucid.xlsx")
        self.assertEqual(second.output.name, "campana_Lucid_2.xlsx")
        self.assertEqual(default_output(source, self.folder, source_system="lucid").name,
                         "campana_Lucid_3.xlsx")
        self.assertEqual(first.output.read_bytes(), before)
        self.assertEqual(self.read_output(first.output), [("0990000001", "PERSONA", "", "BUZON")])
        with self.assertRaises(BaseError):
            convert_files([source], first.output, source_system="lucid")
        self.assertEqual(first.output.read_bytes(), before)

    def test_output_collision_retries_without_overwriting(self):
        source = self.csv_file([["0990000001", "PERSONA", "VENTA"]])
        calls = []

        def racing_writer(output, *args, **kwargs):
            calls.append(output)
            if len(calls) == 1:
                output.write_bytes(b"existing workbook")
            return write_lucid_workbook(output, *args, **kwargs)

        with patch("scripts.transformar_base.write_lucid_workbook", side_effect=racing_writer):
            result = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual(calls[0].read_bytes(), b"existing workbook")
        self.assertEqual(result.output.name, "campana_Lucid_2.xlsx")
        self.assertEqual(self.read_output(result.output), [("0990000001", "PERSONA", "", "VENTA")])

    def test_consolidation_deduplicates_across_files_after_normalizing(self):
        first = self.csv_file([["0990000001", "PRIMERO", "VENTA"]])
        second = self.csv_file([["593990000001", "SEGUNDO", "NO CONTESTA"],
                                ["990000002", "TERCERO", "BUZON"]], filename="otro.csv")
        result = convert_files([first, second], output_folder=self.folder, source_system="lucid")
        self.assertEqual((result.filtered, result.unique), (3, 2))
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "PRIMERO", "", "VENTA"), ("0990000002", "TERCERO", "", "BUZON"),
        ])

    def test_company_indicators_are_explicit_and_do_not_match_people_or_ruc(self):
        companies = ["SERVICIOS S.A.", "INDUSTRIAS S.A.S.", "SERVICIOS s. a. s.",
                     "COMPAÑÍA LTDA.", "ABC CIA. LIMITADA", "ABC COMPAÑÍA LIMITADA",
                     "SERVICIOS S.C.C.", "SOLUCIONES SAS", "SOCIEDAD ANÓNIMA EJEMPLO",
                     "ABC SOCIEDAD POR ACCIONES SIMPLIFICADA", "ABC S.R.L.", "ABC E.I.R.L."]
        people = ["SARAH PEREZ", "SASHA GODOY", "SANTIAGO SALAS", "MARIA SA",
                  "ANA MARIA S", "S.A.SALUD", "ABCS.A.", "PEDRO COMPANIA",
                  "NOMBRE COMERCIAL AMBIGUO", "", "JUAN PEREZ 1791234567001"]
        for name in companies:
            with self.subTest(name=name):
                self.assertTrue(is_company_name(name))
        for name in people:
            with self.subTest(name=name):
                self.assertFalse(is_company_name(name))

    def test_default_company_exclusion_happens_before_dedup_and_preserves_person_ruc(self):
        source = self.csv_file([
            ["990000001", "EMPRESA S.A.", "1790000000001", "EMPRESA"],
            ["0990000001", "SARAH PEREZ", "0012345678001", "PERSONA"],
            ["990000002", "SASHA GODOY", "0001234567", "PERSONA DOS"],
            ["990000003", "OTRA COMPAÑÍA LTDA.", "1790000000002", "EMPRESA DOS"],
        ], headers=["Celular", "Nombre", "Identificación", "Estado"])
        original = source.read_bytes()
        result = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual((result.filtered, result.unique, result.companies, result.excluded), (2, 2, 2, 2))
        self.assertEqual(result.entity_filter, "people")
        self.assertEqual(result.sources, [("campana", 4, 2)])
        self.assertEqual(self.read_output(result.output), [
            ("0990000001", "SARAH PEREZ", "0012345678001", "PERSONA"),
            ("0990000002", "SASHA GODOY", "0001234567", "PERSONA DOS"),
        ])
        self.assertEqual(source.read_bytes(), original)
        all_result = convert_files([source], output_folder=self.folder, source_system="lucid", entity_filter="all")
        self.assertEqual((all_result.filtered, all_result.unique, all_result.companies, all_result.excluded), (4, 3, 2, 0))
        companies_result = convert_files([source], output_folder=self.folder, source_system="lucid", entity_filter="companies")
        self.assertEqual((companies_result.filtered, companies_result.unique, companies_result.excluded), (2, 2, 2))
        self.assertTrue(all(is_company_name(row[1]) for row in self.read_output(companies_result.output)))

    def test_numeric_formatted_id_and_formula_like_id_remain_text(self):
        path = self.folder / "ids.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["Celular", "Nombre", "ID", "Estado"])
        sheet.append([990000001, "PERSONA UNO", 12345678, "Estado"])
        sheet["C2"].number_format = "0000000000"
        sheet.append([990000002, "PERSONA DOS", "0012345678001", "Estado"])
        book.save(path)
        book.close()
        result = convert_files([path], output_folder=self.folder, source_system="lucid")
        self.assertEqual([row[2] for row in self.read_output(result.output)], ["0012345678", "0012345678001"])
        source = self.csv_file([["990000003", "PERSONA", "=1+1", "PENDIENTE"]],
                               headers=["Celular", "Nombre", "ID", "Estado"])
        result = convert_files([source], output_folder=self.folder, source_system="lucid")
        book = openpyxl.load_workbook(result.output, data_only=False)
        try:
            self.assertEqual(book["Hoja1"]["C2"].value, "=1+1")
            self.assertEqual(book["Hoja1"]["C2"].data_type, "s")
        finally:
            book.close()

    def test_unknown_entity_filter_fails_before_output(self):
        source = self.csv_file([["990000001", "PERSONA", "VENTA"]])
        with self.assertRaisesRegex(BaseError, "Filtro de entidades"):
            convert_files([source], output_folder=self.folder, source_system="lucid", entity_filter="guess")
        self.assertFalse(list(self.folder.glob("*.xlsx")))

    def test_new_four_column_workbook_can_be_converted_again_without_losing_id(self):
        source = self.csv_file([["990000001", "PERSONA", "0012345678", "PENDIENTE"]],
                               headers=["Celular", "Nombre", "ID", "Estado"])
        first = convert_files([source], output_folder=self.folder, source_system="lucid")
        second = convert_files([first.output], output_folder=self.folder, source_system="lucid")
        self.assertEqual(self.read_output(first.output), self.read_output(second.output))

    def test_all_companies_produces_empty_four_column_workbook_without_changing_issabel(self):
        source = self.csv_file([["990000001", "EMPRESA S.A.", "1790000000001", "PENDIENTE"]],
                               headers=["Celular", "Nombre", "ID", "Estado"])
        result = convert_files([source], output_folder=self.folder, source_system="lucid")
        self.assertEqual((result.companies, result.excluded, result.filtered, result.unique), (1, 1, 0, 0))
        self.assertEqual(self.read_output(result.output), [])
        source = self.csv_file([
            ["0990000001", "Success", "Agente", "2026-09-16", "30", "EMPRESA S.A.",
             "0123456789", "ELIMINAR DE BASE DE DATOS"],
        ], headers=["Teléfono", "Estado Llamada", "Agente", "Fecha y hora", "Duración(Seg)",
                    "Nombre", "ID", "ESTADO"], filename="issabel.csv")
        result = convert_files([source], output_folder=self.folder, entity_filter="people")
        self.assertEqual((result.source_system, result.filtered, result.unique, result.excluded), ("issabel", 1, 1, 0))


if __name__ == "__main__":
    unittest.main()
